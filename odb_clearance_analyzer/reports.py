"""Report exporters for ODB++ clearance analysis."""

from __future__ import annotations

import csv
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from .models import (
    AnalysisResult,
    FeatureAttributeRecord,
    EffectiveAirGapRecord,
    GeometryDebugRecord,
    MeasurementRecord,
    PerNetMinimum,
    ProgressCallback,
)
from .settings_profile import DEFAULT_SETTINGS_FILENAME, analysis_config_to_profile, write_settings_profile
from .voltage_estimator import material_group_from_cti, voltage_settings_summary, ipc2221a_settings_text
from .voltage_guessing.assignment_store import load_assignment_store
from .voltage_guessing.models import AssignmentStore
from .voltage_guessing.normalization import normalize_net_name_safe
from .voltage_guessing.requirements import (
    VoltageRequirementResolver,
    VoltageRequirementResult,
    summarize_voltage_requirements_for_measurements,
    voltage_requirement_csv_fields,
    voltage_requirement_excel_headers,
    voltage_requirement_excel_values,
    voltage_standard_compliance_csv_fields,
    voltage_standard_compliance_excel_headers,
    voltage_standard_compliance_excel_values,
    voltage_standard_compliance_row,
)


REPORT_SETTINGS_APP_VERSION = "0.4.37"


class ReportWriter:
    """Write analysis results to CSV, Markdown, Excel and ZIP package."""

    def __init__(self, progress: ProgressCallback | None = None):
        self.progress = progress or (lambda _message: None)
        self._assignment_store_cache: dict[Path, AssignmentStore | None] = {}
        # One VoltageRequirementResolver per AnalysisResult; the strong result
        # reference in the value keeps id(result) stable for the cache key.
        self._voltage_resolver_cache: dict[int, tuple[AnalysisResult, VoltageRequirementResolver]] = {}

    def write_all(self, result: AnalysisResult) -> dict[str, Path]:
        out = result.config.output_dir
        out.mkdir(parents=True, exist_ok=True)
        files = {
            "full_csv": out / "net_to_net_measurements_full.csv",
            "critical_csv": out / f"net_to_net_critical_under_{self._threshold_name(result.config.threshold_mm)}mm.csv",
            "per_net_csv": out / "net_to_net_per_net_minimum.csv",
            "feature_attributes_csv": out / "odb_feature_attributes.csv",
            "geometry_debug_csv": out / "geometry_debug_critical_pairs.csv",
            "assigned_voltages_csv": out / "net_voltage_assignments.csv",
            "settings_json": out / DEFAULT_SETTINGS_FILENAME,
            "markdown": out / "net_to_net_clearance_report.md",
            "excel": out / "net_to_net_clearance_report.xlsx",
        }

        self._write_measurements_csv(files["full_csv"], result.measurements, result)
        self._write_measurements_csv(files["critical_csv"], result.critical_measurements, result)
        self._write_per_net_csv(files["per_net_csv"], result.per_net_minimum, result)
        self._write_feature_attributes_csv(
            files["feature_attributes_csv"],
            list(result.job.iter_feature_attributes(include_empty=False)),
        )
        self._write_geometry_debug_csv(files["geometry_debug_csv"], result.geometry_debug_records)
        self._write_assigned_voltages_csv(files["assigned_voltages_csv"], result)
        settings_profile = analysis_config_to_profile(
            result.config,
            app_version=REPORT_SETTINGS_APP_VERSION,
            include_paths=True,
            source="automatic_report_export",
        )
        write_settings_profile(files["settings_json"], settings_profile)
        if result.effective_air_gap_records:
            files["effective_air_gap_csv"] = out / "net_to_net_effective_air_gap_matrix.csv"
            self._write_effective_air_gap_csv(files["effective_air_gap_csv"], result.effective_air_gap_records, result)
        self._write_markdown(files["markdown"], result)
        self._write_excel(files["excel"], result)

        archive_path = out / "net_to_net_report_package.zip"
        with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for file_path in files.values():
                if file_path.exists():
                    zf.write(file_path, arcname=file_path.name)
        files["zip"] = archive_path
        result.report_files.update(files)
        return files

    def _export_effective_voltage(self, result_or_config) -> bool:
        """Whether to include CTI/pollution/altitude voltage estimate in exported files."""

        config = getattr(result_or_config, "config", result_or_config)
        return bool(getattr(config, "export_effective_max_voltage", True))

    def _export_ipc_voltage(self, result_or_config) -> bool:
        """Whether to include IPC-2221A/B voltage estimate in exported files."""

        config = getattr(result_or_config, "config", result_or_config)
        return bool(getattr(config, "export_ipc2221a_max_voltage", True))

    def _voltage_csv_fields(self, result_or_config) -> list[str]:
        fields: list[str] = []
        if self._export_effective_voltage(result_or_config):
            fields.append("effective_max_voltage_v")
        if self._export_ipc_voltage(result_or_config):
            fields.append("ipc2221a_max_voltage_v")
        return fields

    def _voltage_requirement_fields(self) -> list[str]:
        return voltage_requirement_csv_fields()

    def _assignment_store_for_result(self, result: AnalysisResult) -> AssignmentStore | None:
        settings = dict(getattr(result.config, "voltage_guessing", {}) or {})
        raw_path = str(settings.get("assignment_store_path", "net_voltage_assignments.json") or "net_voltage_assignments.json")
        path = Path(raw_path)
        if not path.is_absolute():
            path = Path(result.config.output_dir) / path
        if path in self._assignment_store_cache:
            return self._assignment_store_cache[path]
        try:
            store = load_assignment_store(path) if path.exists() else None
        except Exception:
            store = None
        self._assignment_store_cache[path] = store
        return store

    def _voltage_resolver_for_result(self, result: AnalysisResult) -> VoltageRequirementResolver:
        cached = self._voltage_resolver_cache.get(id(result))
        if cached is not None and cached[0] is result:
            return cached[1]
        store = self._assignment_store_for_result(result)
        settings = dict(getattr(result.config, "voltage_guessing", {}) or {})
        if store is not None:
            merged_settings = dict(getattr(store, "settings", {}) or {})
            merged_settings.update(settings)
            resolver = VoltageRequirementResolver(store.assignments, merged_settings)
        else:
            resolver = VoltageRequirementResolver({}, settings)
        self._voltage_resolver_cache[id(result)] = (result, resolver)
        return resolver

    def _voltage_requirement_for_pair(self, result: AnalysisResult, net_a: str, net_b: str) -> VoltageRequirementResult:
        return self._voltage_resolver_for_result(result).resolve(net_a, net_b)

    def _voltage_requirement_row(self, result: AnalysisResult, net_a: str, net_b: str) -> dict[str, object]:
        return self._voltage_requirement_for_pair(result, net_a, net_b).as_row()

    def _voltage_standard_fields(self) -> list[str]:
        return voltage_standard_compliance_csv_fields()

    def _voltage_standard_row(self, record, req: VoltageRequirementResult) -> dict[str, object]:
        # OK/NOK is intentionally based on the IEC-style effective max voltage
        # estimate. The actual voltage is the hierarchy-resolved required voltage,
        # so galvanic-zone priority is respected.
        return voltage_standard_compliance_row(req.required_voltage_v, getattr(record, "effective_max_voltage_v", None))

    def _voltage_standard_excel_values(self, record, req: VoltageRequirementResult) -> list[object]:
        return voltage_standard_compliance_excel_values(req.required_voltage_v, getattr(record, "effective_max_voltage_v", None))

    def _assignment_store_settings_for_result(self, result: AnalysisResult) -> dict[str, object]:
        store = self._assignment_store_for_result(result)
        settings = dict(getattr(result.config, "voltage_guessing", {}) or {})
        if store is not None:
            merged = dict(getattr(store, "settings", {}) or {})
            merged.update(settings)
            return merged
        return settings

    def _assigned_voltage_rows(self, result: AnalysisResult) -> list[dict[str, object]]:
        """Return report-ready voltage-assignment rows sorted by net name."""

        store = self._assignment_store_for_result(result)
        if store is None or not getattr(store, "assignments", None):
            return []
        rows: list[dict[str, object]] = []
        for assignment in sorted(store.assignments.values(), key=lambda item: item.net_name):
            evidence = getattr(assignment, "evidence", None)
            waiver = getattr(assignment, "waiver", None)
            rows.append({
                "net_name": assignment.net_name,
                "normalized_name": normalize_net_name_safe(assignment.net_name),
                "assigned_voltage_v": "" if assignment.final_voltage_v is None else assignment.final_voltage_v,
                "final_class": assignment.final_class,
                "reference_net": assignment.reference_net,
                "voltage_type": assignment.voltage_type,
                "source": assignment.source,
                "confidence": assignment.confidence,
                "severity": assignment.severity,
                "review_state": assignment.review_state,
                "galvanic_zone": getattr(assignment, "galvanic_zone", ""),
                "winning_rule_id": getattr(evidence, "matched_rule_id", "") if evidence is not None else "",
                "all_matched_rule_ids": ";".join(getattr(evidence, "all_matched_rule_ids", []) or []) if evidence is not None else "",
                "rule_reason": getattr(evidence, "rule_reason", "") if evidence is not None else "",
                "rule_warning": getattr(evidence, "rule_warning", "") if evidence is not None else "",
                "reviewed_by": assignment.reviewed_by,
                "reviewed_at_utc": assignment.reviewed_at_utc,
                "review_reason": assignment.review_reason,
                "waived": "true" if waiver is not None and waiver.waived else "false",
                "waiver_reason": getattr(waiver, "waiver_reason", "") if waiver is not None else "",
                "waiver_scope": getattr(waiver, "waiver_scope", "") if waiver is not None else "",
                "notes": assignment.notes,
                "source_revision": assignment.source_revision,
                "last_seen_revision": assignment.last_seen_revision,
            })
        return rows

    def _assigned_voltage_fields(self) -> list[str]:
        return [
            "net_name",
            "normalized_name",
            "assigned_voltage_v",
            "final_class",
            "reference_net",
            "voltage_type",
            "source",
            "confidence",
            "severity",
            "review_state",
            "galvanic_zone",
            "winning_rule_id",
            "all_matched_rule_ids",
            "rule_reason",
            "rule_warning",
            "reviewed_by",
            "reviewed_at_utc",
            "review_reason",
            "waived",
            "waiver_reason",
            "waiver_scope",
            "notes",
            "source_revision",
            "last_seen_revision",
        ]

    def _write_assigned_voltages_csv(self, path: Path, result: AnalysisResult) -> int:
        self.progress(f"Writing {path.name}")
        rows = self._assigned_voltage_rows(result)
        with path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=self._assigned_voltage_fields(), extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        return len(rows)

    def _voltage_excel_headers(self, result_or_config) -> list[str]:
        headers: list[str] = []
        if self._export_effective_voltage(result_or_config):
            headers.append("Effective max voltage V")
        if self._export_ipc_voltage(result_or_config):
            headers.append("IPC-2221A max voltage V")
        return headers

    def _voltage_excel_values(self, record, result_or_config) -> list[object]:
        values: list[object] = []
        if self._export_effective_voltage(result_or_config):
            values.append(record.effective_max_voltage_v)
        if self._export_ipc_voltage(result_or_config):
            values.append(record.ipc2221a_max_voltage_v)
        return values

    def _markdown_voltage_headers(self, result_or_config) -> tuple[list[str], list[str]]:
        headers: list[str] = []
        aligns: list[str] = []
        if self._export_effective_voltage(result_or_config):
            headers.append("Effective max voltage V")
            aligns.append("---:")
        if self._export_ipc_voltage(result_or_config):
            headers.append("IPC-2221A max voltage V")
            aligns.append("---:")
        return headers, aligns

    def _markdown_voltage_cells(self, record, result_or_config) -> list[str]:
        cells: list[str] = []
        if self._export_effective_voltage(result_or_config):
            cells.append("" if record.effective_max_voltage_v is None else f"{record.effective_max_voltage_v:.1f}")
        if self._export_ipc_voltage(result_or_config):
            cells.append("" if record.ipc2221a_max_voltage_v is None else f"{record.ipc2221a_max_voltage_v:.1f}")
        return cells

    def _markdown_requirement_cells(self, req: VoltageRequirementResult, record=None) -> list[str]:
        def fmt(value):
            return "" if value is None else f"{float(value):.1f}"
        status, margin = ("", None)
        if record is not None:
            values = self._voltage_standard_excel_values(record, req)
            status, margin = values[0], values[1]
        return [
            fmt(req.required_voltage_v),
            status,
            fmt(margin),
            req.requirement_source,
            req.zone_a,
            req.zone_b,
            fmt(req.local_voltage_v),
            fmt(req.voltage_difference_v),
            req.warning.replace("|", "/"),
        ]

    def _write_measurements_csv(self, path: Path, records: list[MeasurementRecord], result: AnalysisResult) -> None:
        self.progress(f"Writing {path.name}")
        fieldnames = [
            "layer",
            "net_a",
            "net_b",
            "clearance_mm",
            *self._voltage_csv_fields(result),
            *self._voltage_requirement_fields(),
            *self._voltage_standard_fields(),
            "x_a_mm",
            "y_a_mm",
            "x_b_mm",
            "y_b_mm",
        ]
        with path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for record in records:
                row = record.as_row()
                req = self._voltage_requirement_for_pair(result, record.net_a, record.net_b)
                row.update(req.as_row())
                row.update(self._voltage_standard_row(record, req))
                writer.writerow(row)

    def _write_per_net_csv(self, path: Path, records: list[PerNetMinimum], result: AnalysisResult) -> None:
        self.progress(f"Writing {path.name}")
        fieldnames = [
            "net",
            "min_clearance_mm",
            *self._voltage_csv_fields(result),
            *self._voltage_requirement_fields(),
            *self._voltage_standard_fields(),
            "layer",
            "other_net",
            "x_this_mm",
            "y_this_mm",
            "x_other_mm",
            "y_other_mm",
        ]
        with path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for record in records:
                row = record.as_row()
                req = self._voltage_requirement_for_pair(result, record.net, record.other_net)
                row.update(req.as_row())
                row.update(self._voltage_standard_row(record, req))
                writer.writerow(row)

    def _write_feature_attributes_csv(
        self, path: Path, records: list[FeatureAttributeRecord]
    ) -> None:
        self.progress(f"Writing {path.name}")
        fieldnames = [
            "layer",
            "feature_index",
            "net",
            "feature_kind",
            "polarity",
            "symbol_name",
            "raw_attributes",
            "decoded_attributes",
            "raw_feature",
        ]
        with path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for record in records:
                writer.writerow(record.as_row())

    def _write_geometry_debug_csv(
        self, path: Path, records: list[GeometryDebugRecord]
    ) -> None:
        self.progress(f"Writing {path.name}")
        fieldnames = [
            "layer",
            "net_a",
            "net_b",
            "clearance_mm",
            "x_a_mm",
            "y_a_mm",
            "x_b_mm",
            "y_b_mm",
            "relation",
            "intersection_area_mm2",
            "intersection_length_mm",
            "geom_a_type",
            "geom_b_type",
            "geom_a_bounds_mm",
            "geom_b_bounds_mm",
            "zero_clearance_kind",
            "feature_hits_a",
            "feature_hits_b",
            "negative_hits_near_point",
            "notes",
        ]
        with path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for record in records:
                writer.writerow(record.as_row())


    def _write_effective_air_gap_csv(
        self, path: Path, records: list[EffectiveAirGapRecord], result: AnalysisResult
    ) -> None:
        self.progress(f"Writing {path.name}")
        fieldnames = [
            "layer",
            "net_a",
            "net_b",
            "direct_clearance_mm",
            "effective_air_gap_mm",
            *self._voltage_csv_fields(result),
            *self._voltage_requirement_fields(),
            *self._voltage_standard_fields(),
            "copper_blocked_length_mm",
            "copper_on_path",
            "blocker_nets",
            "x_a_mm",
            "y_a_mm",
            "x_b_mm",
            "y_b_mm",
        ]
        with path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for record in records:
                row = record.as_row()
                req = self._voltage_requirement_for_pair(result, record.net_a, record.net_b)
                row.update(req.as_row())
                row.update(self._voltage_standard_row(record, req))
                writer.writerow(row)

    def _write_markdown(self, path: Path, result: AnalysisResult) -> None:
        self.progress(f"Writing {path.name}")
        minimum = result.minimum_clearance_mm
        minimum_text = "N/A" if minimum is None else f"{minimum:.6f} mm"
        source = result.job.source_path.name if result.job.source_path else str(result.job.root)
        now = datetime.now().isoformat(timespec="seconds")
        feature_attr_rows = list(result.job.iter_feature_attributes(include_empty=False))
        lines = [
            "# ODB++ Net-to-Net Copper Clearance Report",
            "",
            f"Generated: **{now}**",
            "",
            "## Input",
            "",
            f"- ODB++ source: `{source}`",
            f"- Step: `{result.job.step_name}`",
            f"- Signal layers: `{', '.join(result.job.signal_layers)}`",
            f"- PCB outline source: **{getattr(result.job, 'pcb_outline_source', None) or 'Not found'}**",
            f"- Critical threshold: **{result.config.threshold_mm:.3f} mm**",
            f"- Include `$NONE$` net: **{result.config.include_none_net}**",
            f"- Effective Net-to-Net distance enabled: **{result.config.effective_air_gap_matrix}**",
            f"- IEC 60664-1 settings: **{voltage_settings_summary(result.config.cti, result.config.pollution_degree, result.config.altitude_m)}**",
            f"- IEC 60664-1 layer pollution degrees: **{', '.join(f'{layer}=PD{pd}' for layer, pd in result.config.layer_pollution_degrees.items()) or 'auto'}**",
            f"- External conformal coating: **{result.config.external_conformal_coating}**",
            f"- Metallic particle size: **{result.config.metallic_particle_size_mm:g} mm**",
            f"- Layer roles for IPC-2221A/B estimate: **{', '.join(f'{layer}={role}' for layer, role in result.config.layer_roles.items()) or 'auto'}**",
            f"- Export Effective max voltage V: **{self._export_effective_voltage(result)}**",
            f"- Export IPC-2221A max voltage V: **{self._export_ipc_voltage(result)}**",
            f"- Settings JSON: **{DEFAULT_SETTINGS_FILENAME}**",
            "",
            "## Summary",
            "",
            f"- Nets in EDA data: **{len(result.job.nets_by_number)}**",
            f"- Named nets in analysis: **{result.named_net_count}**",
            f"- Measured layer/net pairs: **{result.measurement_count}**",
            f"- Critical pairs below threshold: **{result.critical_count}**",
            f"- Minimum same-layer copper-to-copper clearance: **{minimum_text}**",
            f"- ODB++ feature attribute rows decoded: **{len(feature_attr_rows)}**",
            f"- Geometry debug rows: **{len(result.geometry_debug_records)}**",
            f"- Effective Net-to-Net distance rows: **{len(result.effective_air_gap_records)}**",
            "",
            "## Layer Net Counts",
            "",
            "| Layer | Nets measured | Skipped/unsupported features |",
            "|---|---:|---:|",
        ]
        for layer in result.job.signal_layers:
            lines.append(
                f"| {layer} | {result.layer_net_counts.get(layer, 0)} | {result.skipped_features.get(layer, 0)} |"
            )

        voltage_headers, voltage_aligns = self._markdown_voltage_headers(result)
        req_headers = ["Required V", "OK/NOK", "Margin V", "Source", "Zone A", "Zone B", "Local V", "ΔV", "Warning"]
        req_aligns = ["---:", "---", "---:", "---", "---", "---", "---:", "---:", "---"]
        critical_headers = ["Rank", "Layer", "Net A", "Net B", "Clearance mm", *voltage_headers, *req_headers, "Point A mm", "Point B mm"]
        critical_aligns = ["---:", "---", "---", "---", "---:", *voltage_aligns, *req_aligns, "---", "---"]
        lines.extend(["", "## Most Critical Measurements", "", "| " + " | ".join(critical_headers) + " |", "| " + " | ".join(critical_aligns) + " |"])
        for rank, record in enumerate(result.measurements[:50], start=1):
            cells = [
                str(rank),
                record.layer,
                f"`{record.net_a}`",
                f"`{record.net_b}`",
                f"{record.clearance_mm:.6f}",
                *self._markdown_voltage_cells(record, result),
                *self._markdown_requirement_cells(self._voltage_requirement_for_pair(result, record.net_a, record.net_b), record),
                f"({record.x_a_mm or 0.0:.3f}, {record.y_a_mm or 0.0:.3f})",
                f"({record.x_b_mm or 0.0:.3f}, {record.y_b_mm or 0.0:.3f})",
            ]
            lines.append("| " + " | ".join(cells) + " |")

        lines.extend(
            [
                "",
                "## Geometry Debug for Critical Pairs",
                "",
                "This section is intended to diagnose suspicious `0 mm` rows. `Point A` and `Point B` are generated closest-point coordinates from the reconstructed geometry; they are not necessarily original track endpoints.",
                "",
                "| Layer | Net A | Net B | Clearance mm | Relation | Zero kind | Intersection area mm² | Source hits A | Source hits B | Negative/cutouts near point | Notes |",
                "|---|---|---|---:|---|---|---:|---|---|---|---|",
            ]
        )
        for record in result.geometry_debug_records[:50]:
            row = record.as_row()
            lines.append(
                f"| {record.layer} | `{record.net_a}` | `{record.net_b}` | {record.clearance_mm:.6f} | "
                f"{record.relation} | {record.zero_clearance_kind} | "
                f"{'' if record.intersection_area_mm2 is None else f'{record.intersection_area_mm2:.9g}'} | "
                f"`{row['feature_hits_a']}` | `{row['feature_hits_b']}` | "
                f"`{row['negative_hits_near_point']}` | {record.notes} |"
            )
        if len(result.geometry_debug_records) > 50:
            lines.append(f"| ... | ... | ... | ... | ... | ... | ... | ... | ... | ... | ... {len(result.geometry_debug_records) - 50} more rows in CSV/Excel |")

        if result.effective_air_gap_records:
            voltage_headers, voltage_aligns = self._markdown_voltage_headers(result)
            req_headers = ["Required V", "OK/NOK", "Margin V", "Source", "Zone A", "Zone B", "Local V", "ΔV", "Warning"]
            req_aligns = ["---:", "---", "---:", "---", "---", "---", "---:", "---:", "---"]
            effective_headers = [
                "Rank",
                "Layer",
                "Net A",
                "Net B",
                "Direct clearance mm",
                "Effective Net-to-Net distance mm",
                *voltage_headers,
                *req_headers,
                "Copper blocked mm",
                "Blocker nets",
            ]
            effective_aligns = ["---:", "---", "---", "---", "---:", "---:", *voltage_aligns, *req_aligns, "---:", "---"]
            lines.extend(
                [
                    "",
                    "## Effective Net-to-Net distance",
                    "",
                    "This optional diagnostic uses the closest-point line between two nets and subtracts copper occupied by any other net on that line. It is useful when intermediate copper splits the line-of-sight path, but it is still not a routed creepage solver.",
                    "",
                    "| " + " | ".join(effective_headers) + " |",
                    "| " + " | ".join(effective_aligns) + " |",
                ]
            )
            for rank, record in enumerate(result.effective_air_gap_records[:50], start=1):
                cells = [
                    str(rank),
                    record.layer,
                    f"`{record.net_a}`",
                    f"`{record.net_b}`",
                    f"{record.direct_clearance_mm:.6f}",
                    f"{record.effective_air_gap_mm:.6f}",
                    *self._markdown_voltage_cells(record, result),
                    *self._markdown_requirement_cells(self._voltage_requirement_for_pair(result, record.net_a, record.net_b), record),
                    f"{record.copper_blocked_length_mm:.6f}",
                    f"`{record.blocker_nets}`",
                ]
                lines.append("| " + " | ".join(cells) + " |")
            if len(result.effective_air_gap_records) > 50:
                lines.append(f"| ... | ... | ... | ... | ... | ... | ... | ... {len(result.effective_air_gap_records) - 50} more rows in CSV/Excel |")

        lines.extend(
            [
                "",
                "## ODB++ Feature Attribute Decoding",
                "",
                "ODB++ feature attribute suffixes such as `0=5,3=0.3` are now decoded using the layer feature-file `@` attribute-name table and `&` text-string table.",
                "The original raw attribute text is still retained for traceability.",
                "",
                "| Layer | Feature | Net | Kind | Raw attributes | Decoded attributes |",
                "|---|---:|---|---|---|---|",
            ]
        )
        for record in feature_attr_rows[:50]:
            lines.append(
                f"| {record.layer} | {record.feature_index} | `{record.net}` | {record.feature_kind} | "
                f"`{record.raw_attributes}` | `{record.decoded_attributes}` |"
            )
        if len(feature_attr_rows) > 50:
            lines.append(f"| ... | ... | ... | ... | ... | ... {len(feature_attr_rows) - 50} more rows in CSV/Excel |")

        assigned_rows = self._assigned_voltage_rows(result)
        lines.extend([
            "",
            "## Assigned Voltages",
            "",
            "This table is exported completely in `net_voltage_assignments.csv`. The preview below shows the first 100 assigned nets used by the voltage-requirement resolver.",
            "",
            "| Net | Assigned V | Class | Source | Review | Zone | Notes |",
            "|---|---:|---|---|---|---|---|",
        ])
        for row in assigned_rows[:100]:
            voltage = row["assigned_voltage_v"]
            voltage_text = "" if voltage == "" else f"{float(voltage):.1f}"
            notes = str(row.get("notes", "") or "").replace("|", "/")
            lines.append(
                f"| `{row['net_name']}` | {voltage_text} | {row['final_class']} | {row['source']} | "
                f"{row['review_state']} | {row['galvanic_zone']} | {notes} |"
            )
        if len(assigned_rows) > 100:
            lines.append(f"| ... | ... | ... | ... | ... | ... | ... {len(assigned_rows) - 100} more rows in CSV/Excel |")

        hierarchy_summary = summarize_voltage_requirements_for_measurements(result.measurements, self._assignment_store_for_result(result) or {}, getattr(self._assignment_store_for_result(result), "settings", getattr(result.config, "voltage_guessing", {})))
        lines.extend(
            [
                "",
                "## Voltage Requirement Hierarchy",
                "",
                "Priority: different galvanic zones use the configured Zone 1 ↔ Zone 2 working voltage; same-zone or incomplete-zone pairs use local net/manual/class voltage.",
                "",
                f"- Zone 1 ↔ Zone 2 working voltage: **{float(hierarchy_summary['zone_voltage_v'] or 0):g} V**",
                f"- Zone 1 nets: **{hierarchy_summary['zone_1_nets']}**",
                f"- Zone 2 nets: **{hierarchy_summary['zone_2_nets']}**",
                f"- Cross-zone measured pairs: **{hierarchy_summary['cross_zone_pairs']}**",
                f"- Cross-zone PASS / FAIL / UNKNOWN: **{hierarchy_summary['cross_zone_pass']} / {hierarchy_summary['cross_zone_fail']} / {hierarchy_summary['cross_zone_unknown']}**",
                f"- Pairs using local net voltage: **{hierarchy_summary['local_voltage_pairs']}**",
                f"- Pairs with incomplete-zone warning: **{hierarchy_summary['incomplete_zone_pairs']}**",
            ]
        )

        lines.extend(
            [
                "",
                "## Notes and Limitations",
                "",
                "- This report measures **same-layer copper-to-copper spacing** from ODB++ copper features.",
                *(
                    ["- `Effective max voltage V` is a conservative IEC 60664-1-style screening estimate from measured spacing, CTI/material group, per-layer pollution degree, altitude correction, and external-layer metallic-particle correction; it is not a formal IEC/UL certification result."]
                    if self._export_effective_voltage(result)
                    else []
                ),
                *(
                    ["- `IPC-2221A max voltage V` is an IPC-2221A/B-style conductor-spacing screening estimate using the selected layer role (`internal`/`external`) and altitude class."]
                    if self._export_ipc_voltage(result)
                    else []
                ),
                "- `OK/NOK` compares the hierarchy-resolved `Required voltage V` against `Effective max voltage V`; OK means required voltage is less than or equal to the calculated supported voltage, NOK means it is higher.",
                "- It does not automatically prove IEC/UL compliance. A voltage/net-class table and safety engineer review are required for formal pass/fail rules.",
                "- Rectangular rounded pads are measured with a conservative rectangular approximation.",
                "- Internal creepage inside connectors, relays, optocouplers, cable assemblies, and component bodies is not available from PCB ODB++ copper data alone.",
                "- Solder mask is not treated as reliable insulation unless the project explicitly qualifies it as such.",
            ]
        )
        if result.warnings:
            lines.extend(["", "## Parser Warnings", ""])
            for warning in result.warnings[:100]:
                lines.append(f"- {warning}")
            if len(result.warnings) > 100:
                lines.append(f"- ... {len(result.warnings) - 100} more warnings omitted")

        path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    def _write_excel(self, path: Path, result: AnalysisResult) -> None:
        self.progress(f"Writing {path.name}")
        # write_only mode keeps large jobs responsive and avoids very high memory use.
        wb = Workbook(write_only=True)

        ws = wb.create_sheet("Summary")
        self._append_rows(
            ws,
            [
                ["Field", "Value"],
                ["ODB++ source", result.job.source_path.name if result.job.source_path else str(result.job.root)],
                ["Step", result.job.step_name],
                ["Signal layers", ", ".join(result.job.signal_layers)],
                ["PCB outline source", getattr(result.job, "pcb_outline_source", None) or "Not found"],
                ["Threshold mm", result.config.threshold_mm],
                ["Include $NONE$", result.config.include_none_net],
                ["Effective Net-to-Net distance enabled", result.config.effective_air_gap_matrix],
                ["CTI", result.config.cti],
                ["Material group", material_group_from_cti(result.config.cti)],
                ["Base/external pollution degree", result.config.pollution_degree],
                ["IEC layer pollution degrees", ", ".join(f"{layer}=PD{pd}" for layer, pd in result.config.layer_pollution_degrees.items())],
                ["External conformal coating", result.config.external_conformal_coating],
                ["Metallic particle size mm", result.config.metallic_particle_size_mm],
                ["Altitude m", result.config.altitude_m],
                ["IPC layer roles", ", ".join(f"{layer}={role}" for layer, role in result.config.layer_roles.items())],
                ["Export Effective max voltage V", self._export_effective_voltage(result)],
                ["Export IPC-2221A max voltage V", self._export_ipc_voltage(result)],
                ["Nets in EDA data", len(result.job.nets_by_number)],
                ["Named nets in analysis", result.named_net_count],
                ["Measurements", result.measurement_count],
                ["Critical below threshold", result.critical_count],
                ["Minimum clearance mm", result.minimum_clearance_mm if result.minimum_clearance_mm is not None else "N/A"],
                ["Feature attribute rows decoded", result.feature_attribute_count],
                ["Geometry debug rows", len(result.geometry_debug_records)],
                ["Effective Net-to-Net distance rows", len(result.effective_air_gap_records)],
                [],
                ["Layer", "Nets measured", "Skipped/unsupported features"],
                *[
                    [layer, result.layer_net_counts.get(layer, 0), result.skipped_features.get(layer, 0)]
                    for layer in result.job.signal_layers
                ],
            ],
        )

        self._append_measurement_sheet(wb, "Critical", result.critical_measurements, result)
        self._append_measurement_sheet(wb, "All measurements", result.measurements, result)
        self._append_per_net_sheet(wb, result.per_net_minimum, result)
        if result.effective_air_gap_records:
            self._append_effective_air_gap_sheet(wb, result.effective_air_gap_records, result)
        self._append_geometry_debug_sheet(wb, result.geometry_debug_records)
        self._append_feature_attributes_sheet(
            wb, list(result.job.iter_feature_attributes(include_empty=False))
        )
        self._append_assigned_voltages_sheet(wb, result)
        ws_warn = wb.create_sheet("Warnings")
        self._append_rows(ws_warn, [["Warning"], *[[w] for w in result.warnings]])
        wb.save(path)

    def _append_rows(self, ws, rows: list[list[object]]) -> None:
        for row in rows:
            ws.append(row)

    def _append_measurement_sheet(
        self, wb: Workbook, title: str, records: list[MeasurementRecord], result: AnalysisResult
    ) -> None:
        ws = wb.create_sheet(title[:31])
        ws.append([
            "Layer",
            "Net A",
            "Net B",
            "Clearance mm",
            *self._voltage_excel_headers(result),
            *voltage_requirement_excel_headers(),
            *voltage_standard_compliance_excel_headers(),
            "Point A X mm",
            "Point A Y mm",
            "Point B X mm",
            "Point B Y mm",
        ])
        for rec in records:
            ws.append([
                rec.layer,
                rec.net_a,
                rec.net_b,
                rec.clearance_mm,
                *self._voltage_excel_values(rec, result),
                *voltage_requirement_excel_values((req := self._voltage_requirement_for_pair(result, rec.net_a, rec.net_b))),
                *self._voltage_standard_excel_values(rec, req),
                rec.x_a_mm,
                rec.y_a_mm,
                rec.x_b_mm,
                rec.y_b_mm,
            ])

    def _append_per_net_sheet(self, wb: Workbook, records: list[PerNetMinimum], result: AnalysisResult) -> None:
        ws = wb.create_sheet("Per net minimum")
        ws.append([
            "Net",
            "Minimum clearance mm",
            *self._voltage_excel_headers(result),
            *voltage_requirement_excel_headers(),
            *voltage_standard_compliance_excel_headers(),
            "Layer",
            "Other net",
            "This point X mm",
            "This point Y mm",
            "Other point X mm",
            "Other point Y mm",
        ])
        for rec in records:
            ws.append([
                rec.net,
                rec.min_clearance_mm,
                *self._voltage_excel_values(rec, result),
                *voltage_requirement_excel_values((req := self._voltage_requirement_for_pair(result, rec.net, rec.other_net))),
                *self._voltage_standard_excel_values(rec, req),
                rec.layer,
                rec.other_net,
                rec.x_this_mm,
                rec.y_this_mm,
                rec.x_other_mm,
                rec.y_other_mm,
            ])

    def _append_effective_air_gap_sheet(
        self, wb: Workbook, records: list[EffectiveAirGapRecord], result: AnalysisResult
    ) -> None:
        ws = wb.create_sheet("Effective Net-to-Net distance")
        ws.append([
            "Layer",
            "Net A",
            "Net B",
            "Direct clearance mm",
            "Effective Net-to-Net distance mm",
            *self._voltage_excel_headers(result),
            *voltage_requirement_excel_headers(),
            *voltage_standard_compliance_excel_headers(),
            "Copper blocked length mm",
            "Copper on path",
            "Blocker nets",
            "Point A X mm",
            "Point A Y mm",
            "Point B X mm",
            "Point B Y mm",
        ])
        for rec in records:
            ws.append([
                rec.layer,
                rec.net_a,
                rec.net_b,
                rec.direct_clearance_mm,
                rec.effective_air_gap_mm,
                *self._voltage_excel_values(rec, result),
                *voltage_requirement_excel_values((req := self._voltage_requirement_for_pair(result, rec.net_a, rec.net_b))),
                *self._voltage_standard_excel_values(rec, req),
                rec.copper_blocked_length_mm,
                rec.copper_on_path,
                rec.blocker_nets,
                rec.x_a_mm,
                rec.y_a_mm,
                rec.x_b_mm,
                rec.y_b_mm,
            ])


    def _append_assigned_voltages_sheet(self, wb: Workbook, result: AnalysisResult) -> None:
        ws = wb.create_sheet("Assigned voltages")
        fields = self._assigned_voltage_fields()
        headers = [field.replace("_", " ").title().replace("V", "V") for field in fields]
        ws.append(headers)
        for row in self._assigned_voltage_rows(result):
            ws.append([row.get(field, "") for field in fields])


    def _append_geometry_debug_sheet(
        self, wb: Workbook, records: list[GeometryDebugRecord]
    ) -> None:
        ws = wb.create_sheet("Geometry debug")
        headers = [
            "Layer",
            "Net A",
            "Net B",
            "Clearance mm",
            "Point A X mm",
            "Point A Y mm",
            "Point B X mm",
            "Point B Y mm",
            "Relation",
            "Intersection area mm2",
            "Intersection length mm",
            "Geom A type",
            "Geom B type",
            "Geom A bounds mm",
            "Geom B bounds mm",
            "Zero clearance kind",
            "Feature hits A",
            "Feature hits B",
            "Negative/cutout hits near point",
            "Notes",
        ]
        ws.append(headers)
        for rec in records:
            row = rec.as_row()
            ws.append([
                row["layer"],
                row["net_a"],
                row["net_b"],
                row["clearance_mm"],
                row["x_a_mm"],
                row["y_a_mm"],
                row["x_b_mm"],
                row["y_b_mm"],
                row["relation"],
                row["intersection_area_mm2"],
                row["intersection_length_mm"],
                row["geom_a_type"],
                row["geom_b_type"],
                row["geom_a_bounds_mm"],
                row["geom_b_bounds_mm"],
                row["zero_clearance_kind"],
                row["feature_hits_a"],
                row["feature_hits_b"],
                row["negative_hits_near_point"],
                row["notes"],
            ])

    def _append_feature_attributes_sheet(
        self, wb: Workbook, records: list[FeatureAttributeRecord]
    ) -> None:
        ws = wb.create_sheet("Feature attributes")
        ws.append([
            "Layer",
            "Feature index",
            "Net",
            "Feature kind",
            "Polarity",
            "Symbol name",
            "Raw attributes",
            "Decoded attributes",
            "Raw feature",
        ])
        for rec in records:
            ws.append([
                rec.layer,
                rec.feature_index,
                rec.net,
                rec.feature_kind,
                rec.polarity,
                rec.symbol_name,
                rec.raw_attributes,
                rec.decoded_attributes,
                rec.raw_feature,
            ])

    def _threshold_name(self, threshold: float) -> str:
        return str(threshold).replace(".", "p")
